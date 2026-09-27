"""Sample files for the live-ingestion part of the demo. They are NOT loaded into the
database: the presenter drops them into the Ingest screen during the meeting."""
import copy
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..config import INBOX_DIR
from ..fmt import fmt_number
from ..utils import add_months, last_day
from .documents import write_change_order, write_schedule_xml
from .simulate import simulate
from .specs import PROJECTS

SAMPLES = [
    dict(
        filename="Historico_Linea_Flujo_Tiputini_2019.xlsx", kind="legacy_workbook",
        es="Liquidación final de un proyecto de 2019 en un formato antiguo y desordenado: celdas combinadas, números guardados como texto y nombres de partidas abreviados.",
        en="Final cost settlement of a 2019 project in an old, messy format: merged cells, numbers stored as text and shortened line names.",
        pt="Liquidação final de um projeto de 2019 em formato antigo e desorganizado: células mescladas, números guardados como texto e nomes de itens abreviados.",
    ),
    dict(
        filename="EC-2503_OC-012_Variante_km48.pdf", kind="change_order_pdf",
        es="Nueva orden de cambio (número 12) del Oleoducto Amazonía – Tramo B, en PDF.",
        en="New change order (number 12) for the Amazon Pipeline – Section B, as a PDF.",
        pt="Nova ordem de alteração (número 12) do Oleoduto Amazônia – Trecho B, em PDF.",
    ),
    dict(
        filename="EC-2503_OC-013_escaneada.pdf", kind="scanned_pdf",
        es="Otra orden de cambio (número 13) del mismo proyecto, pero escaneada: es solo una imagen, sin texto.",
        en="Another change order (number 13) for the same project, but scanned: it is only an image, with no text.",
        pt="Outra ordem de alteração (número 13) do mesmo projeto, mas digitalizada: é só uma imagem, sem texto.",
    ),
    dict(
        filename="PE-2410_Cronograma_rev5.xml", kind="schedule_xml",
        es="Revisión 5 del cronograma del Ducto Selva Central, exportada desde Microsoft Project.",
        en="Revision 5 of the Central Jungle pipeline schedule, exported from Microsoft Project.",
        pt="Revisão 5 do cronograma do Duto Selva Central, exportada do Microsoft Project.",
    ),
]

LEGACY_ROWS = [
    # label, weight, final-cost multiplier, note
    ("Movilizac./campamento", 0.06, 1.05, ""),
    ("Apertura DDV", 0.09, 1.12, ""),
    ("Excav. zanja", 0.11, 1.08, ""),
    ("Tendido/curvado", 0.07, 1.02, ""),
    ("Sold. + END", 0.21, 1.04, "Incluye radiografía"),
    ("Revest. juntas", 0.05, 1.00, ""),
    ("Bajado/tapado", 0.10, 1.22, "Lluvias feb–abr 2020"),
    ("Cruces rio (PHD)", 0.09, 1.35, "Cruce adicional río Tiputini (no reconocido)"),
    ("Prueba hidrost.", 0.04, 1.00, ""),
    ("Restauracion", 0.06, 1.03, ""),
    ("Indirectos/GG", 0.12, 1.15, "Ampliación de plazo 3 meses"),
]
LEGACY_CONTRACT = 12_480_000.0


def build_legacy_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws.merge_cells("A1:F1")
    ws["A1"] = "CONDUTO S.A. – LIQUIDACIÓN FINAL DE OBRA"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A3"] = "PROYECTO: LÍNEA DE FLUJO TIPUTINI – FASE 1"
    ws["A4"] = "CLIENTE: Oriente Crudos S.A."
    ws["A5"] = "UBICACIÓN: Oriente – Orellana (Amazonía)"
    ws["A6"] = "DIÁMETRO: 10 pulg."
    ws["D6"] = "LONGITUD: 21,4 km"
    ws["A7"] = "INICIO: mar-2019"
    ws["D7"] = "FIN: jun-2020"
    ws["A8"] = f"MONTO CONTRATO (USD): {fmt_number(LEGACY_CONTRACT, 'es')}"
    for c, h in enumerate(["ITEM", "DESCRIPCION", "PPTO OFERTA", "COSTO FINAL", "DESV.", "OBS."], start=1):
        cell = ws.cell(row=10, column=c, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
    budget_total = LEGACY_CONTRACT * 0.842
    tb = ta = 0.0
    for i, (label, w, mult, note) in enumerate(LEGACY_ROWS, start=1):
        r = 10 + i
        budget = round(budget_total * w, 2)
        actual = round(budget * mult, 2)
        tb += budget
        ta += actual
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=label)
        # messy on purpose: some amounts typed as text with Spanish separators
        ws.cell(row=r, column=3, value=fmt_number(budget, "es") if i % 3 == 0 else budget)
        ws.cell(row=r, column=4, value=fmt_number(actual, "es") if i % 4 == 0 else actual)
        ws.cell(row=r, column=5, value=round(actual - budget, 2))
        ws.cell(row=r, column=6, value=note or None)
    r = 11 + len(LEGACY_ROWS)
    ws.cell(row=r, column=2, value="TOTAL").font = Font(bold=True)
    ws.cell(row=r, column=3, value=round(tb, 2))
    ws.cell(row=r, column=4, value=round(ta, 2))
    ws.cell(row=r, column=5, value=round(ta - tb, 2))
    ws.cell(row=r + 2, column=1, value="Elaborado por: Dpto. de Costos – Datos sintéticos para demostración")
    ws.column_dimensions["B"].width = 26
    for col in "CDE":
        ws.column_dimensions[col].width = 16
    ws.column_dimensions["F"].width = 40
    wb.save(path)


def build_scanned_change_order(path):
    """Image-only PDF, like a desk scan: no text layer at all."""
    import reportlab
    from PIL import Image, ImageDraw, ImageFilter, ImageFont
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    fonts = Path(reportlab.__file__).parent / "fonts"
    regular = ImageFont.truetype(str(fonts / "Vera.ttf"), 30)
    bold = ImageFont.truetype(str(fonts / "VeraBd.ttf"), 40)
    small = ImageFont.truetype(str(fonts / "Vera.ttf"), 24)
    img = Image.new("L", (1654, 2339), 246)
    d = ImageDraw.Draw(img)
    d.rectangle([110, 110, 1544, 200], fill=60)
    d.text((140, 128), "CONDUTO", font=bold, fill=250)
    d.text((1120, 140), "Orden de cambio OC-013", font=small, fill=235)
    d.text((140, 280), "ORDEN DE CAMBIO N° OC-013", font=bold, fill=30)
    rows = [
        ("Proyecto", "EC-2503 – Oleoducto Amazonía – Tramo B"), ("Contrato", "EC-2503-C"), ("Contratante", "Oriente Crudos S.A."),
        ("Fecha de presentación", "22/09/2026"), ("Actividad afectada", "Bajado y tapado"), ("Causa", "Solicitud del cliente"),
        ("Monto solicitado", "USD 415.300,00"), ("Impacto en plazo", "12 días"), ("Estado", "Pendiente de aprobación"),
    ]
    y = 390
    for k, v in rows:
        d.rectangle([140, y - 8, 620, y + 40], fill=228)
        d.text((155, y), k, font=regular, fill=35)
        d.text((660, y), v, font=regular, fill=25)
        y += 70
    d.text((140, y + 30), "Descripción", font=bold, fill=30)
    desc = ["Protección adicional de tubería con sacos de suelo-cemento en el cruce",
            "de quebradas entre el km 52 y el km 55, solicitada por la fiscalización."]
    for i, line in enumerate(desc):
        d.text((140, y + 100 + i * 44), line, font=regular, fill=30)
    for x in (140, 900):
        d.line([x, 1900, x + 520, 1900], fill=40, width=3)
    d.text((140, 1915), "Por el Contratista", font=small, fill=40)
    d.text((900, 1915), "Por el Contratante", font=small, fill=40)
    d.ellipse([1150, 1650, 1450, 1850], outline=90, width=6)
    d.text((1195, 1730), "RECIBIDO", font=bold, fill=90)
    d.text((140, 2240), "DEMO – DATOS SINTÉTICOS", font=small, fill=120)
    rng = random.Random(13)
    for _ in range(2500):  # scanner grain
        x, yy = rng.randrange(img.width), rng.randrange(img.height)
        d.point((x, yy), fill=rng.randrange(150, 230))
    img = img.rotate(0.7, fillcolor=240, resample=Image.BICUBIC).filter(ImageFilter.GaussianBlur(0.7))
    c = canvas.Canvas(str(path), pagesize=A4)
    c.drawImage(ImageReader(img.convert("RGB")), 0, 0, width=A4[0], height=A4[1])
    c.save()


def build_inbox(fx: dict):
    for f in INBOX_DIR.glob("*"):
        f.unlink()
    build_legacy_workbook(INBOX_DIR / "Historico_Linea_Flujo_Tiputini_2019.xlsx")
    build_scanned_change_order(INBOX_DIR / "EC-2503_OC-013_escaneada.pdf")

    spec = next(p for p in PROJECTS if p["code"] == "EC-2503")
    sim = simulate(spec, fx)
    co = dict(
        number="OC-012", cause="geotech", activity="DDV", amount=float(round(sim.contract * 0.011, -2)),
        status="pending", submitted_date="2026-09-15", decision_date=None, period="2026-09",
        schedule_impact_days=20, executed=1.0, budget_split={},
        titles=(
            "Variante del trazado entre el km 48+100 y el km 48+900 por deslizamiento activo, incluida la estabilización del talud",
            "Route variant between km 48+100 and km 48+900 due to an active landslide, including slope stabilization",
            "Variante do traçado entre o km 48+100 e o km 48+900 por deslizamento ativo, incluindo estabilização do talude",
        ),
    )
    write_change_order(sim, co, INBOX_DIR / "EC-2503_OC-012_Variante_km48.pdf")

    spec = next(p for p in PROJECTS if p["code"] == "PE-2410")
    sim = simulate(spec, fx)
    sim2 = copy.copy(sim)
    sim2.tasks = []
    for t in sim.tasks:
        t = dict(t)
        if t["pct"] < 1:
            t["finish"] = last_day(add_months(t["finish"][:7], 2))
        sim2.tasks.append(t)
    sim2.A = sim.A + 2
    write_schedule_xml(sim2, INBOX_DIR / "PE-2410_Cronograma_rev5.xml", name_override=f"{spec['code']} {spec['name']} – rev5")
