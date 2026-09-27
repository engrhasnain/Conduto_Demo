"""Microsoft Dynamics GP "SmartList" exports: general-ledger transactions with
project-segmented cost accounts, one workbook per country.

GP books cost by account (labor, equipment, materials, subcontracts), not by field
activity, so it reconciles against the Excel control sheets at project x month x
cost-type level. Three discrepancies are planted on purpose, the kind every
contractor has: an invoice posted in GP but missing in Excel, an Excel accrual not
yet in GP, and a transaction duplicated in GP."""
import random

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..config import STATUS_PERIOD
from ..i18n import COST_TYPE_ORDER
from ..utils import add_months, last_day
from .simulate import SimResult

WINDOW_START = add_months(STATUS_PERIOD, -11)  # a 12-month fiscal export

ACCOUNT_BASE = {"MO": "5101-01", "EQ": "5102-01", "MAT": "5103-01", "SUB": "5104-01"}
ACCOUNT_DESC = {
    "es": {"MO": "Costo mano de obra directa", "EQ": "Alquiler y operación de equipos", "MAT": "Materiales de construcción", "SUB": "Subcontratos de obra"},
    "pt": {"MO": "Custo mão de obra direta", "EQ": "Aluguel e operação de equipamentos", "MAT": "Materiais de construção", "SUB": "Subempreitadas de obra"},
}
HEADERS = {
    "es": ["Número de diario", "Fecha de TRX", "Número de cuenta", "Descripción de cuenta", "Proyecto", "Importe del débito", "Importe del crédito", "Referencia", "Origen", "Id. de proveedor"],
    "pt": ["Número do diário", "Data da TRX", "Número da conta", "Descrição da conta", "Projeto", "Valor do débito", "Valor do crédito", "Referência", "Origem", "ID do fornecedor"],
}
VENDORS = {
    "EC": {"EQ": ["MAQORIENTE", "EQUIPESADO", "GRUASNAPO"], "MAT": ["FERRETECUADOR", "HORMIGONESQ", "REVESTIPLAST"], "SUB": ["PERFORIENTE", "RADIOINSP", "TOPOANDES"]},
    "PE": {"EQ": ["MAQSELVA", "ALQUIPERU", "GRUASCENTRO"], "MAT": ["ACEROSLIMA", "CONCRETOSAC", "TUBOREVEST"], "SUB": ["PHDPERU", "ENDSERVIC", "SERVGEOTEC"]},
    "BR": {"EQ": ["LOCAMAQ", "EQUIPNORTE", "GUINDASTES"], "MAT": ["ACOBRASIL", "CONCRETNE", "REVESTSUL"], "SUB": ["PERFURANE", "INSPECTEND", "TOPONORDES"]},
}
SOURCES = {"MO": "UPRCC", "EQ": "PMTRX", "MAT": "PMTRX", "SUB": "PMTRX"}
SPLITS = {"MO": (1, 2), "EQ": (2, 4), "MAT": (2, 4), "SUB": (1, 3)}

# (project, period, cost type) -> kind of discrepancy
PLANTED = {
    ("EC-2503", "2026-08", "SUB"): ("missing_in_excel", 185_000.0, "FAC-004471", "PERFORIENTE"),
    ("PE-2410", "2026-06", "EQ"): ("missing_in_erp", 240_000.0, None, None),
    ("BR-2501", "2026-07", "MAT"): ("duplicate", None, None, None),
}


def _split(total: float, k: int, rng: random.Random) -> list[float]:
    weights = [rng.uniform(0.6, 1.4) for _ in range(k)]
    s = sum(weights)
    parts = [round(total * w / s, 2) for w in weights]
    parts[-1] = round(total - sum(parts[:-1]), 2)
    return parts


def write_gp_export(country: str, sims: list[SimResult], path, lang: str) -> list[dict]:
    """Writes the SmartList workbook and returns the transactions with their row locators."""
    rng = random.Random(f"gp-{country}")
    wb = Workbook()
    ws = wb.active
    ws.title = "SmartList"
    for c, h in enumerate(HEADERS[lang], start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    rows: list[dict] = []
    journal = {"EC": 184_200, "PE": 96_400, "BR": 57_100}[country]

    for sim in sims:
        code = sim.spec["code"]
        for t, period in enumerate(sim.periods):
            if period < WINDOW_START:
                continue
            for ct in COST_TYPE_ORDER:
                total = sum(sim.costs[a][ct][t] for a in sim.costs)
                if total <= 0:
                    continue
                planted = PLANTED.get((code, period, ct))
                gp_total = total - (planted[1] if planted and planted[0] == "missing_in_erp" else 0.0)
                k = rng.randint(*SPLITS[ct])
                parts = _split(gp_total, k, rng)
                for i, amount in enumerate(parts):
                    day = min(5 + i * 7 + rng.randint(0, 5), 28)
                    if ct == "MO":
                        ref, vendor = f"NOM-{period.replace('-', '')}-{i + 1}", None
                    else:
                        vendor = rng.choice(VENDORS[country][ct])
                        ref = f"FAC-{rng.randint(1000, 9999):06d}"
                    rows.append(dict(project=code, period=period, trx_date=f"{period}-{day:02d}", cost_type=ct, debit=amount, credit=0.0,
                                     reference=ref, source=SOURCES[ct], vendor=vendor))
                if planted and planted[0] == "missing_in_excel":
                    rows.append(dict(project=code, period=period, trx_date=last_day(period), cost_type=ct, debit=planted[1], credit=0.0,
                                     reference=planted[2], source="PMTRX", vendor=planted[3]))
                if planted and planted[0] == "duplicate":
                    dup = dict(rows[-1])
                    rows.append(dup)

    rows.sort(key=lambda r: (r["trx_date"], r["project"], r["cost_type"]))
    for i, r in enumerate(rows, start=2):
        journal += 1
        r["journal"] = str(journal)
        r["account"] = f"{ACCOUNT_BASE[r['cost_type']]}-{r['project'].replace('-', '')}"
        values = [r["journal"], r["trx_date"], r["account"], ACCOUNT_DESC[lang][r["cost_type"]], r["project"],
                  r["debit"], r["credit"] or None, r["reference"], r["source"], r["vendor"]]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=i, column=c, value=v)
            if c in (6, 7):
                cell.number_format = "#,##0.00"
        r["locator"] = f"SmartList!F{i}"
    for col, width in zip("ABCDEFGHIJ", (14, 12, 20, 34, 12, 16, 16, 18, 9, 16)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    wb.save(path)
    return rows
