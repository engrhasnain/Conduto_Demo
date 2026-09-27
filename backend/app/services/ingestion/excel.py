"""Format-agnostic parser for project cost-control workbooks.

Detects the header row and the role of every column (code, activity, cost type,
budget, month...), classifies each sheet (monthly costs, physical progress,
production, summary, legacy close-out table), keeps a cell locator for every value
and reconciles the parsed numbers against the file's own total rows."""
import hashlib
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .normalize import norm, parse_number, parse_period

HEADER_ROLES = [
    ("cost_type", ["tipo de costo", "recurso", "natureza", "cost type", "resource", "tipo costo"]),
    ("code", ["cod", "codigo", "item", "partida", "wbs", "n", "no", "nro"]),
    ("activity", ["actividad", "descripcion", "atividade", "descricao", "activity", "description", "concepto", "rubro"]),
    ("budget", ["presupuesto", "ppto", "orcamento", "budget", "presup"]),
    ("actual", ["costo final", "costo real", "custo final", "real acumulado", "realizado acumulado", "actual cost", "gasto real"]),
    ("variance", ["desv", "desviacion", "variacion", "variance", "diferencia"]),
    ("notes", ["obs", "observaciones", "notas", "comentarios", "notes"]),
    ("total", ["total"]),
    ("series", ["tipo", "prog real", "type", "serie"]),
]

SERIES_VALUES = {
    "planned": ["programado", "prog", "previsto", "planned", "plan", "planificado"],
    "actual": ["real", "realizado", "actual", "ejecutado"],
}

META_KEYS = [
    ("contract_value", ["monto de contrato", "monto contratado", "monto del contrato", "monto contrato", "valor do contrato", "contract value"]),
    ("bid_margin", ["margen de oferta", "margen ofertado", "margem da proposta", "bid margin"]),
    ("project", ["proyecto", "projeto", "project", "obra"]),
    ("client", ["cliente", "client", "contratante"]),
    ("currency", ["moneda", "moeda", "currency"]),
    ("cutoff", ["fecha de corte", "data de corte", "corte"]),
    ("location", ["ubicacion", "localizacao", "location", "region"]),
    ("diameter", ["diametro", "diameter"]),
    ("length", ["longitud", "comprimento", "length", "extension"]),
    ("start", ["fecha de inicio", "inicio", "start"]),
    ("finish", ["fecha de fin", "fin", "termino", "finish", "cierre"]),
    ("terrain", ["terreno", "zona", "terrain"]),
]

PROJECT_CODE_RE = re.compile(r"\b([A-Z]{2}-\d{4})\b")


def header_role(h) -> str | None:
    n = norm(h)
    if not n:
        return None
    for role, kws in HEADER_ROLES:
        for kw in kws:
            if n == kw or n.startswith(kw + " ") or (len(kw) >= 5 and n.startswith(kw)):
                return role
    return None


def series_of(v) -> str | None:
    n = norm(v)
    for key, vals in SERIES_VALUES.items():
        if n in vals:
            return key
    return None


def _grid(ws) -> list[list]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _loc(sheet: str, r: int, c: int) -> str:
    return f"{sheet}!{get_column_letter(c + 1)}{r + 1}"


def _find_monthly_header(grid) -> int | None:
    best = None
    for r, row in enumerate(grid[:30]):
        months = sum(1 for v in row if parse_period(v))
        texts = sum(1 for v in row if isinstance(v, str) and v.strip() and not parse_period(v))
        if months >= 2 and texts >= 1:
            score = months * 10 + texts
            if best is None or score > best[1]:
                best = (r, score)
    return best[0] if best else None


def _find_table_header(grid) -> int | None:
    for r, row in enumerate(grid[:40]):
        roles = {header_role(v) for v in row if isinstance(v, str)}
        if {"activity", "budget"} <= roles and ("actual" in roles or "code" in roles):
            return r
    return None


def _is_production(grid) -> int | None:
    for r, row in enumerate(grid[:12]):
        labels = [norm(v) for v in row if isinstance(v, str)]
        if any(l.startswith("km") for l in labels) and any(("junta" in l or "weld" in l or "repar" in l or "repair" in l) for l in labels):
            return r
    return None


def extract_meta(grids: dict[str, list[list]]) -> dict:
    meta = {}
    for sheet, grid in grids.items():
        for r, row in enumerate(grid[:14]):
            cells = [(c, v) for c, v in enumerate(row) if v not in (None, "")]
            for i, (c, v) in enumerate(cells):
                if not isinstance(v, str) or ":" not in v:
                    continue
                key_text, value, vc = v, None, c
                if ":" in v and v.split(":", 1)[1].strip():
                    key_text, value = v.split(":", 1)
                    value = value.strip()
                elif i + 1 < len(cells):
                    vc, value = cells[i + 1]
                nk = norm(key_text)
                for field, kws in META_KEYS:
                    if field in meta:
                        continue
                    if any(nk == kw or nk.startswith(kw) for kw in kws) and value not in (None, ""):
                        meta[field] = {"value": value, "label": key_text.strip().rstrip(":"), "locator": _loc(sheet, r, vc)}
                        break
    out = {}
    for field, item in meta.items():
        raw = item["value"]
        val = raw
        if field in ("contract_value", "length", "diameter"):
            val = parse_number(raw)
        elif field == "bid_margin":
            x = parse_number(raw)
            val = x / 100 if x is not None and x > 1 else x
        elif field in ("start", "finish", "cutoff"):
            val = parse_period(raw) or raw
        else:
            val = str(raw).strip()
        out[field] = {"value": val, "raw": str(raw), "label": item["label"], "locator": item["locator"]}
    if "project" in out:
        m = PROJECT_CODE_RE.search(str(out["project"]["raw"]))
        if m:
            out["project_code"] = {"value": m.group(1), "raw": m.group(1), "label": out["project"]["label"], "locator": out["project"]["locator"]}
    if "currency" not in out:
        blob = " ".join(str(v) for g in grids.values() for row in g[:14] for v in row if isinstance(v, str))
        for cur, pat in (("USD", r"\bUSD\b|US\$"), ("PEN", r"S/\.?"), ("BRL", r"R\$")):
            if re.search(pat, blob):
                out["currency"] = {"value": cur, "raw": cur, "label": "currency", "locator": None}
                break
    else:
        v = str(out["currency"]["value"]).upper()
        out["currency"]["value"] = "PEN" if "S/" in v or "SOL" in v else "BRL" if "R$" in v or "REA" in v else "USD" if "US" in v or "DOL" in v else v
    return out


def parse_workbook(path) -> dict:
    wb = load_workbook(path, data_only=True)
    grids = {ws.title: _grid(ws) for ws in wb.worksheets}
    result = dict(kind="unknown", sheets=[], cost_rows=[], progress_rows=[], production_rows=[],
                  legacy_rows=[], skipped=[], checks=[], periods=[], fingerprint=None, headers_signature=None)
    result["meta"] = extract_meta(grids)
    periods = set()
    signature_parts = []
    table_candidates = []  # close-out tables only count when there are no monthly cost sheets

    for sheet, grid in grids.items():
        prod_hr = _is_production(grid)
        if prod_hr is not None:
            header = grid[prod_hr]
            cols = {}
            for c, v in enumerate(header):
                n = norm(v)
                if n.startswith("km"):
                    cols["km"] = c
                elif "junta" in n or "weld" in n:
                    cols["welds"] = c
                elif "repar" in n or "repair" in n:
                    cols["repairs"] = c
            for r in range(prod_hr + 1, len(grid)):
                row = grid[r]
                p = parse_period(row[0]) if row else None
                if not p:
                    continue
                periods.add(p)
                result["production_rows"].append(dict(
                    period=p, km=parse_number(row[cols["km"]]) if "km" in cols else None,
                    welds=parse_number(row[cols["welds"]]) if "welds" in cols else None,
                    repairs=parse_number(row[cols["repairs"]]) if "repairs" in cols else None,
                    locator=_loc(sheet, r, cols.get("km", 1)),
                ))
            result["sheets"].append(dict(name=sheet, role="production", header_row=prod_hr + 1, columns=[
                dict(letter=get_column_letter(c + 1), header=str(header[c]), role=role, confidence=0.9) for role, c in cols.items()]))
            continue

        hr = _find_monthly_header(grid)
        if hr is not None:
            header = grid[hr]
            columns, month_cols = [], {}
            role_cols = {}
            for c, v in enumerate(header):
                p = parse_period(v)
                if p:
                    month_cols[c] = p
                    columns.append(dict(letter=get_column_letter(c + 1), header=p, role="month", confidence=1.0))
                    continue
                role = header_role(v)
                if role and role not in role_cols:
                    role_cols[role] = c
                if v not in (None, ""):
                    columns.append(dict(letter=get_column_letter(c + 1), header=str(v), role=role or "ignored", confidence=0.95 if role else 0.0))
            if "activity" not in role_cols:
                for c, v in enumerate(header):
                    if c not in month_cols and isinstance(v, str) and c not in role_cols.values():
                        role_cols["activity"] = c
                        break
            # decide series vs cost-type columns from their values
            for role in ("series", "cost_type"):
                if role in role_cols:
                    c = role_cols[role]
                    vals = [grid[r][c] for r in range(hr + 1, min(len(grid), hr + 12)) if c < len(grid[r]) and grid[r][c]]
                    is_series = sum(1 for v in vals if series_of(v)) >= max(1, len(vals) // 2)
                    if role == "cost_type" and is_series:
                        role_cols["series"] = role_cols.pop("cost_type")
                    elif role == "series" and not is_series:
                        role_cols["cost_type"] = role_cols.pop("series")
            for col in columns:
                for role, c in role_cols.items():
                    if col["letter"] == get_column_letter(c + 1):
                        col["role"] = role
            sheet_role = "progress" if "series" in role_cols else "cost"
            periods.update(month_cols.values())
            act_c = role_cols.get("activity")
            total_row = None
            for r in range(hr + 1, len(grid)):
                row = grid[r] + [None] * (len(header) - len(grid[r]))
                label = row[act_c] if act_c is not None else None
                nlabel = norm(label)
                values = {p: parse_number(row[c]) for c, p in month_cols.items()}
                has_numbers = any(v is not None for v in values.values())
                if not nlabel:
                    if has_numbers:
                        code_val = norm(row[role_cols["code"]]) if "code" in role_cols else ""
                        if code_val.startswith("total"):
                            total_row = (r, values, parse_number(row[role_cols["budget"]]) if "budget" in role_cols else None)
                        else:
                            result["skipped"].append(dict(sheet=sheet, row=r + 1, label="", reason="no_label"))
                    continue
                if nlabel.startswith("subtotal") or nlabel.startswith("sub total"):
                    result["skipped"].append(dict(sheet=sheet, row=r + 1, label=str(label), reason="subtotal"))
                    continue
                if nlabel.startswith("total"):
                    total_row = (r, values, parse_number(row[role_cols["budget"]]) if "budget" in role_cols else None)
                    result["skipped"].append(dict(sheet=sheet, row=r + 1, label=str(label), reason="total"))
                    continue
                if not has_numbers:
                    continue
                if sheet_role == "cost":
                    result["cost_rows"].append(dict(
                        sheet=sheet, row=r + 1,
                        code=str(row[role_cols["code"]]) if "code" in role_cols and row[role_cols["code"]] is not None else None,
                        label=str(label).strip(),
                        cost_type_label=str(row[role_cols["cost_type"]]).strip() if "cost_type" in role_cols and row[role_cols["cost_type"]] else None,
                        budget=parse_number(row[role_cols["budget"]]) if "budget" in role_cols else None,
                        budget_loc=_loc(sheet, r, role_cols["budget"]) if "budget" in role_cols else None,
                        values={p: v for p, v in values.items() if v is not None},
                        locs={p: _loc(sheet, r, c) for c, p in month_cols.items() if values[p] is not None},
                    ))
                else:
                    series = series_of(row[role_cols["series"]])
                    if series:
                        result["progress_rows"].append(dict(
                            sheet=sheet, row=r + 1, label=str(label).strip(), series=series,
                            values={p: v for p, v in values.items() if v is not None},
                            locs={p: _loc(sheet, r, c) for c, p in month_cols.items() if values[p] is not None},
                        ))
            if sheet_role == "cost":
                non_month = [norm(h) for h in header if h not in (None, "") and not parse_period(h)]
                signature_parts.append(f"cost:{norm(sheet)}:{','.join(non_month)}")
                if total_row:
                    r, tvals, tbudget = total_row
                    parsed = {p: 0.0 for p in month_cols.values()}
                    for row_ in result["cost_rows"]:
                        if row_["sheet"] == sheet:
                            for p, v in row_["values"].items():
                                parsed[p] += v
                    diffs = [abs((tvals.get(p) or 0) - parsed[p]) for p in parsed]
                    exp_total = sum(v or 0 for v in tvals.values())
                    result["checks"].append(dict(code="total_row", ok=max(diffs, default=0) < 1.0,
                                                 params=dict(sheet=sheet, row=r + 1, expected=exp_total, parsed=sum(parsed.values()))))
                    if tbudget is not None:
                        pb = sum(x["budget"] or 0 for x in result["cost_rows"] if x["sheet"] == sheet)
                        result["checks"].append(dict(code="budget_total", ok=abs(pb - tbudget) < 1.0,
                                                     params=dict(sheet=sheet, expected=tbudget, parsed=pb)))
            result["sheets"].append(dict(name=sheet, role=sheet_role, header_row=hr + 1, columns=columns))
            continue

        thr = _find_table_header(grid)
        if thr is not None:
            header = grid[thr]
            role_cols = {}
            columns = []
            for c, v in enumerate(header):
                role = header_role(v)
                if role and role not in role_cols:
                    role_cols[role] = c
                if v not in (None, ""):
                    columns.append(dict(letter=get_column_letter(c + 1), header=str(v), role=role or "ignored", confidence=0.9 if role else 0.0))
            total = None
            cand = dict(sheet=sheet, rows=[], skipped=[], checks=[], signature=None,
                        info=dict(name=sheet, role="closeout_table", header_row=thr + 1, columns=columns))
            for r in range(thr + 1, len(grid)):
                row = grid[r] + [None] * (len(header) - len(grid[r]))
                label = row[role_cols["activity"]]
                nlabel = norm(label)
                budget = parse_number(row[role_cols["budget"]]) if "budget" in role_cols else None
                actual = parse_number(row[role_cols["actual"]]) if "actual" in role_cols else None
                if not nlabel:
                    continue
                if nlabel.startswith("total"):
                    total = (r, budget, actual)
                    cand["skipped"].append(dict(sheet=sheet, row=r + 1, label=str(label), reason="total"))
                    continue
                if budget is None and actual is None:
                    continue
                cand["rows"].append(dict(
                    sheet=sheet, row=r + 1, label=str(label).strip(),
                    code=str(row[role_cols["code"]]) if "code" in role_cols and row[role_cols["code"]] is not None else None,
                    budget=budget, actual=actual,
                    budget_loc=_loc(sheet, r, role_cols["budget"]) if "budget" in role_cols else None,
                    actual_loc=_loc(sheet, r, role_cols["actual"]) if "actual" in role_cols else None,
                    note=str(row[role_cols["notes"]]).strip() if "notes" in role_cols and row[role_cols["notes"]] else None,
                ))
            non_month = [norm(h) for h in header if h not in (None, "")]
            cand["signature"] = f"table:{norm(sheet)}:{','.join(non_month)}"
            if total:
                r, tb, ta = total
                pb = sum(x["budget"] or 0 for x in cand["rows"])
                pa = sum(x["actual"] or 0 for x in cand["rows"])
                if tb is not None:
                    cand["checks"].append(dict(code="budget_total", ok=abs(pb - tb) < 1.0, params=dict(sheet=sheet, expected=tb, parsed=pb)))
                if ta is not None:
                    cand["checks"].append(dict(code="actual_total", ok=abs(pa - ta) < 1.0, params=dict(sheet=sheet, expected=ta, parsed=pa)))
            table_candidates.append(cand)
            continue

        result["sheets"].append(dict(name=sheet, role="summary", header_row=None, columns=[]))

    for cand in table_candidates:
        if result["cost_rows"]:
            result["sheets"].append(dict(cand["info"], role="summary"))
        else:
            result["legacy_rows"] += cand["rows"]
            result["skipped"] += cand["skipped"]
            result["checks"] += cand["checks"]
            signature_parts.append(cand["signature"])
            result["sheets"].append(cand["info"])

    if result["cost_rows"]:
        result["kind"] = "cost_workbook"
    elif result["legacy_rows"]:
        result["kind"] = "legacy_workbook"
    result["periods"] = sorted(periods)
    if signature_parts:
        sig = "|".join(sorted(signature_parts))
        result["headers_signature"] = sig
        result["fingerprint"] = hashlib.sha1(sig.encode()).hexdigest()
    return result
