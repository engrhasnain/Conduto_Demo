"""Ingestion workflow: analyze a file into a reviewable preview, then import it.

Nothing touches the canonical data until a person confirms the preview. Every
imported value keeps its source locator; every new spreadsheet format becomes a
reusable mapping template."""
import difflib
import logging
import re
import shutil
import uuid
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from ...config import DATA_DIR, UPLOADS_DIR, ai_enabled
from ...i18n import ACTIVITY_NAMES, COUNTRIES, COUNTRY_LANG
from ...models import (BudgetLine, ChangeOrder, CostActual, Document, DocumentChunk, IngestionJob, MappingTemplate,
                       Production, Progress, Project, ProjectKpi, ScheduleTask)
from ..kpis import refresh_kpis
from ...ml import registry
from ...ml.label_model import predict as label_predict
from ..llm import LLMError, extract_json
from ..search import invalidate
from .excel import parse_workbook
from .mapping import map_activity, map_cost_type
from .msp import parse_mspdi
from .normalize import norm
from .pdf import CO_FIELDS, OCR_COST_PER_PAGE_USD, ai_extract_scan, content_state, grounded, read_pdf
from .pdf import extract as pdf_extract

log = logging.getLogger("conduto.ingest")

SMALL_WORDS = {"de", "del", "la", "las", "el", "los", "y", "e", "do", "da", "dos", "das", "of", "the", "and"}
TERRAIN_WORDS = {
    "rainforest": ["amazon", "oriente", "selva", "floresta", "orellana", "sucumbios", "loreto", "napo"],
    "highlands": ["sierra", "andes", "andin", "serra", "altura", "montana"],
    "coast": ["costa", "litoral", "coast", "desierto"],
}


class IngestError(ValueError):
    pass


def _rel(path: Path) -> str:
    return path.relative_to(DATA_DIR).as_posix()


def _store(src: Path, filename: str) -> Path:
    safe = re.sub(r"[^\w.\-]+", "_", Path(filename).name)
    dest = UPLOADS_DIR / f"{uuid.uuid4().hex[:8]}_{safe}"
    shutil.copyfile(src, dest)
    return dest


def _nice_title(s: str) -> str:
    words = s.strip().lower().split()
    return " ".join(w if (i and w in SMALL_WORDS) else w[:1].upper() + w[1:] for i, w in enumerate(words))


def kpi_snapshot(db: Session, project_id: int) -> dict | None:
    k = db.get(ProjectKpi, project_id)
    if not k:
        return None
    return dict(forecast_margin_pct=k.forecast_margin_pct, eac_usd=k.eac_usd, pending_co_usd=k.pending_co_usd,
                pending_co_count=k.pending_co_count, delay_months=k.delay_months, cpi=k.cpi, spi=k.spi,
                progress_pct=k.progress_pct, flags=[f for f in k.flags.split(",") if f], health=k.health)


def match_project(db: Session, code: str | None = None, name: str | None = None) -> dict | None:
    if code:
        p = db.scalar(select(Project).where(Project.code == code))
        if p:
            return dict(code=p.code, name=p.name, currency=p.currency, confidence=0.99, method="code")
    if name:
        best, score = None, 0.0
        for p in db.scalars(select(Project)):
            r = difflib.SequenceMatcher(None, norm(name), norm(p.name)).ratio()
            if r > score:
                best, score = p, r
        if best and score >= 0.8:
            return dict(code=best.code, name=best.name, currency=best.currency, confidence=round(score, 2), method="name")
    return None


def ai_map_labels(labels: list[str]) -> dict:
    schema = {
        "type": "object",
        "properties": {"mappings": {"type": "array", "items": {
            "type": "object",
            "properties": {
                "label": {"type": "string"},
                "activity_code": {"type": "string", "enum": list(ACTIVITY_NAMES) + ["UNMAPPED"]},
                "confidence": {"type": "number"},
                "note": {"type": "string", "description": "Short English note, empty if nothing to flag"},
            },
            "required": ["label", "activity_code", "confidence", "note"],
            "additionalProperties": False,
        }}},
        "required": ["mappings"],
        "additionalProperties": False,
    }
    catalog = "\n".join(f"{c}: {n[0]} | {n[1]} | {n[2]}" for c, n in ACTIVITY_NAMES.items())
    system = ("You map cost-line labels from a pipeline and civil-works contractor's spreadsheets (Spanish or "
              "Portuguese, often abbreviated) to a canonical activity code. If one line combines two activities, "
              "pick the one that normally carries most of the cost and say so in the note. Use UNMAPPED for lines "
              "that are not activities (totals, notes). Confidence is between 0 and 1.")
    content = f"Canonical activities:\n{catalog}\n\nLabels to map:\n" + "\n".join(f"- {l}" for l in labels)
    res = extract_json(system, content, schema, max_tokens=4000)
    return {m["label"]: m for m in res["mappings"]}


def map_labels(labels: list[str], template_map: dict | None) -> tuple[dict, bool]:
    """Known template first, then the locally trained model cross-checked with the rules,
    then Claude for whatever is still uncertain (only when an API key is configured)."""
    out, pending = {}, []
    for label in labels:
        nl = norm(label)
        if template_map and nl in template_map:
            out[label] = dict(code=template_map[nl], confidence=1.0, method="template", note=None, alternatives=[])
        else:
            pending.append(label)
    for label, pred in zip(pending, label_predict(registry.label_model(), pending)):
        rules = map_activity(label)
        code, conf, method = pred["code"], pred["confidence"], "ml"
        if rules["code"] and rules["confidence"] >= 0.95 and rules["code"] != code:
            code, conf, method = rules["code"], 0.9, "rules"  # an exact domain phrase outranks the model
        elif rules["code"] and rules["code"] == code:
            conf = max(conf, rules["confidence"])
        note = None
        if rules["note"] == "combined":
            note, conf = "combined", min(conf, 0.7)
        elif code is None:
            note = "not_activity"
        out[label] = dict(code=code, confidence=round(conf, 2), method=method, note=note,
                          alternatives=[a for a in pred["top"][1:] if a[1] >= 0.05])
    low = [l for l, m in out.items() if m["confidence"] < 0.7]
    ai_used = False
    if low and ai_enabled():
        try:
            ai = ai_map_labels(low)
            for label in low:
                a = ai.get(label)
                if a:
                    out[label] = dict(code=None if a["activity_code"] == "UNMAPPED" else a["activity_code"],
                                      confidence=round(a["confidence"], 2), method="ai", note=a["note"] or None,
                                      alternatives=out[label].get("alternatives", []))
            ai_used = True
        except LLMError as e:
            log.warning("AI mapping failed, keeping rules: %s", e)
    return out, ai_used


def _raw_preview(path: Path, sheet: str | None, header_row: int | None, max_cols: int = 8) -> dict:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
    end = (header_row or 1) + 8
    rows = []
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=end, max_col=max_cols, values_only=True), start=1):
        rows.append(dict(row=r, header=(r == header_row), cells=[None if v is None else (v.strftime("%Y-%m") if hasattr(v, "strftime") else v) for v in row]))
    wb.close()
    return dict(sheet=ws.title, rows=rows)


def _meta_values(meta: dict) -> dict:
    return {k: dict(value=v["value"], label=v["label"], locator=v["locator"]) for k, v in meta.items()}


def _checks_ok(checks: list[dict]) -> bool:
    return all(c["ok"] for c in checks)


# ------------------------------------------------------------------ analyze

def analyze(db: Session, src_path: Path, filename: str) -> IngestionJob:
    stored = _store(src_path, filename)
    ext = stored.suffix.lower()
    try:
        if ext in (".xlsx", ".xlsm"):
            preview = _analyze_excel(db, stored)
        elif ext == ".pdf":
            preview = _analyze_pdf(db, stored)
        elif ext == ".xml":
            preview = _analyze_xml(db, stored)
        else:
            preview = dict(kind="unsupported", error=f"Unsupported file type {ext}")
    except Exception as e:  # surface parser failures in the UI instead of a 500
        log.exception("analyze failed")
        preview = dict(kind="error", error=str(e))
    preview["filename"] = filename
    status = {"unsupported": "failed", "error": "failed", "unknown_workbook": "failed", "scanned_pdf": "needs_ai"}.get(preview["kind"], "preview")
    job = IngestionJob(filename=filename, stored_path=_rel(stored), file_kind=preview["kind"], status=status,
                       ai_used=preview.get("engine") == "ai", preview=preview)
    db.add(job)
    db.commit()
    return job


def _analyze_excel(db: Session, path: Path) -> dict:
    parsed = parse_workbook(path)
    tpl = db.scalar(select(MappingTemplate).where(MappingTemplate.fingerprint == parsed["fingerprint"])) if parsed["fingerprint"] else None
    meta = parsed["meta"]
    main = next((s for s in parsed["sheets"] if s["role"] in ("cost", "closeout_table")), None)
    raw = _raw_preview(path, main["name"] if main else None, main["header_row"] if main else None)
    base = dict(engine="template" if tpl else "ml", ai_available=ai_enabled(), fingerprint=parsed["fingerprint"],
                template=dict(id=tpl.id, name=tpl.name, times_used=tpl.times_used) if tpl else None,
                meta=_meta_values(meta), sheets=parsed["sheets"], skipped=parsed["skipped"][:40], raw_preview=raw)

    if parsed["kind"] == "cost_workbook":
        labels = list(dict.fromkeys([r["label"] for r in parsed["cost_rows"]] + [r["label"] for r in parsed["progress_rows"]]))
        act_map, ai_used = map_labels(labels, tpl.mapping.get("activities") if tpl else None)
        if ai_used:
            base["engine"] = "ai"
        ct_labels = list(dict.fromkeys(r["cost_type_label"] for r in parsed["cost_rows"] if r["cost_type_label"]))
        tmap = tpl.mapping.get("cost_types", {}) if tpl else {}
        ct_map = {}
        for l in ct_labels:
            if norm(l) in tmap:
                ct_map[l] = dict(code=tmap[norm(l)], confidence=1.0, method="template")
            else:
                m = map_cost_type(l)
                ct_map[l] = dict(code=m["code"], confidence=m["confidence"], method=m["method"])
        proj = match_project(db, code=meta.get("project_code", {}).get("value"), name=meta.get("project", {}).get("value"))
        totals = defaultdict(lambda: [0.0, 0.0])
        rows_per_label = defaultdict(int)
        for r in parsed["cost_rows"]:
            code = act_map[r["label"]]["code"] or "?"
            totals[code][0] += r["budget"] or 0
            totals[code][1] += sum(r["values"].values())
            rows_per_label[r["label"]] += 1
        unmapped = [l for l, m in act_map.items() if not m["code"]]
        checks = list(parsed["checks"])
        checks.append(dict(code="rows_mapped", ok=not unmapped, params=dict(mapped=len(labels) - len(unmapped), total=len(labels))))
        checks.append(dict(code="project_identified", ok=proj is not None, params=dict(code=proj["code"] if proj else None)))
        file_cur = meta.get("currency", {}).get("value")
        if proj and file_cur:
            checks.append(dict(code="currency_matches", ok=file_cur == proj["currency"], params=dict(file=file_cur, project=proj["currency"])))
        prog_points = sum(len(r["values"]) for r in parsed["progress_rows"] if r["series"] == "actual")
        base.update(
            kind="cost_workbook", project=proj,
            row_mapping=[dict(label=l, rows=rows_per_label.get(l, 0), **m) for l, m in act_map.items()],
            cost_type_mapping=[dict(label=l, **m) for l, m in ct_map.items()],
            periods=dict(first=parsed["periods"][0] if parsed["periods"] else None, last=parsed["periods"][-1] if parsed["periods"] else None, count=len(parsed["periods"])),
            totals=[dict(activity=a, budget=v[0], actual=v[1]) for a, v in totals.items()],
            records=dict(cost_values=sum(len(r["values"]) for r in parsed["cost_rows"]), budget_lines=sum(1 for r in parsed["cost_rows"] if r["budget"] is not None),
                         progress_points=prog_points, production_points=len(parsed["production_rows"])),
            checks=checks,
        )
        base["ready"] = proj is not None and not unmapped
        return base

    if parsed["kind"] == "legacy_workbook":
        labels = [r["label"] for r in parsed["legacy_rows"]]
        act_map, ai_used = map_labels(labels, tpl.mapping.get("activities") if tpl else None)
        if ai_used:
            base["engine"] = "ai"
        mv = {k: v["value"] for k, v in meta.items()}
        currency = mv.get("currency") or "USD"
        country = {"USD": "EC", "PEN": "PE", "BRL": "BR"}.get(currency, "EC")
        loc_text = norm(f"{mv.get('location', '')} {mv.get('terrain', '')} {mv.get('project', '')}")
        terrain = next((t for t, words in TERRAIN_WORDS.items() if any(w in loc_text for w in words)), "coast")
        start, finish = mv.get("start"), mv.get("finish")
        budget = sum(r["budget"] or 0 for r in parsed["legacy_rows"])
        actual = sum(r["actual"] or 0 for r in parsed["legacy_rows"])
        contract = mv.get("contract_value")
        code = None
        if isinstance(start, str) and re.match(r"\d{4}-\d{2}", start):
            y, m = int(start[:4]), int(start[5:7])
            for _ in range(24):
                cand = f"{country}-{str(y)[2:]}{m:02d}"
                if not db.scalar(select(Project).where(Project.code == cand)):
                    code = cand
                    break
                m = m + 1 if m < 12 else 1
                y = y + (1 if m == 1 else 0)
        name = _nice_title(str(mv.get("project", "Proyecto histórico")))
        existing = match_project(db, name=name)
        new_project = dict(
            code=code, name=name, client=mv.get("client"), country=country, currency=currency, terrain=terrain,
            region=mv.get("location"), kind="pipeline" if mv.get("diameter") else "civil",
            diameter_in=mv.get("diameter"), length_km=mv.get("length"), start_period=start, finish_period=finish,
            contract_value=contract, budget=budget, final_cost=actual,
            bid_margin_pct=(contract - budget) / contract if contract else None,
            final_margin_pct=(contract - actual) / contract if contract else None,
        )
        unmapped = [l for l, m in act_map.items() if not m["code"]]
        required = ["contract_value", "start", "finish"] + (["length", "diameter"] if new_project["kind"] == "pipeline" else [])
        missing = [f for f in required if mv.get(f) in (None, "")]
        checks = list(parsed["checks"])
        checks.append(dict(code="rows_mapped", ok=not unmapped, params=dict(mapped=len(labels) - len(unmapped), total=len(labels))))
        checks.append(dict(code="metadata_complete", ok=not missing, params=dict(missing=missing)))
        checks.append(dict(code="not_duplicate", ok=existing is None, params=dict(code=existing["code"] if existing else None)))
        base.update(
            kind="legacy_workbook", new_project=new_project, existing=existing,
            row_mapping=[dict(label=r["label"], rows=1, budget=r["budget"], actual=r["actual"], note_in_file=r["note"], **act_map[r["label"]]) for r in parsed["legacy_rows"]],
            records=dict(budget_lines=len(parsed["legacy_rows"]), cost_values=len(parsed["legacy_rows"])),
            checks=checks,
        )
        base["ready"] = not unmapped and not missing and existing is None and code is not None
        return base

    base.update(kind="unknown_workbook", error="No cost table or monthly cost sheet was found in this workbook.", checks=parsed["checks"])
    return base


def _analyze_pdf(db: Session, path: Path) -> dict:
    pages = read_pdf(path)
    state = content_state(path, pages)
    if state != "text":
        # No text layer: free extraction is impossible. With an API key the scan goes to Claude's
        # vision; without one it is queued for OCR with a cost estimate (never silently mis-parsed).
        if state == "scanned" and ai_enabled():
            try:
                ai = ai_extract_scan(path)
                conf = ai.pop("confidence", {})
                fields = {f: ai.get(f) for f in CO_FIELDS}
                proj = match_project(db, code=fields.get("project_code"))
                return dict(kind="change_order_pdf" if fields.get("doc_type") == "change_order" else "document_pdf",
                            doc_type=fields.get("doc_type"), engine="ai_vision", ai_available=True, project=proj, content_state=state,
                            fields=[dict(field=f, value=fields.get(f), confidence=round(conf.get(f, 0), 2), grounded=None) for f in fields],
                            pages=len(pages), text_preview="", existing=False,
                            checks=[dict(code="project_identified", ok=proj is not None, params=dict(code=proj["code"] if proj else None))],
                            ready=proj is not None and bool(fields.get("amount")))
            except LLMError as e:
                log.warning("Vision extraction failed: %s", e)
        n = len(pages)
        return dict(kind="scanned_pdf", engine="needs_ai", ai_available=ai_enabled(), content_state=state, pages=n,
                    ocr_estimate_usd=[round(OCR_COST_PER_PAGE_USD[0] * n, 2), round(OCR_COST_PER_PAGE_USD[1] * n, 2)],
                    checks=[dict(code="text_layer", ok=False, params=dict(state=state, pages=n))], ready=False)
    fields, conf, engine = pdf_extract(pages)
    proj = match_project(db, code=fields.get("project_code"))
    kind = "change_order_pdf" if fields.get("doc_type") == "change_order" else "document_pdf"
    existing = None
    if proj and fields.get("co_number"):
        p = db.scalar(select(Project).where(Project.code == proj["code"]))
        existing = db.scalar(select(ChangeOrder).where(ChangeOrder.project_id == p.id, ChangeOrder.number == fields["co_number"]))
    checks = [dict(code="project_identified", ok=proj is not None, params=dict(code=proj["code"] if proj else None))]
    if kind == "change_order_pdf":
        checks.append(dict(code="amount_found", ok=bool(fields.get("amount")), params=dict(amount=fields.get("amount"), currency=fields.get("currency"))))
        checks.append(dict(code="number_found", ok=bool(fields.get("co_number")), params=dict(number=fields.get("co_number"))))
        if proj and fields.get("currency"):
            checks.append(dict(code="currency_matches", ok=fields["currency"] == proj["currency"], params=dict(file=fields["currency"], project=proj["currency"])))
        checks.append(dict(code="not_duplicate", ok=existing is None, params=dict(number=fields.get("co_number"))))
    ground = grounded(fields, pages)
    checks.append(dict(code="grounded", ok=all(ground.values()), params=dict(grounded=sum(ground.values()), total=len(ground))))
    return dict(
        kind=kind, doc_type=fields.get("doc_type"), engine=engine, ai_available=ai_enabled(), project=proj, content_state=state,
        fields=[dict(field=f, value=fields.get(f), confidence=round(conf.get(f, 0), 2), grounded=ground.get(f)) for f in fields],
        pages=len(pages), text_preview=pages[0][:2500] if pages else "", existing=bool(existing), checks=checks,
        ready=proj is not None and (kind != "change_order_pdf" or bool(fields.get("amount"))),
    )


def _analyze_xml(db: Session, path: Path) -> dict:
    s = parse_mspdi(path)
    proj = match_project(db, code=s["project_code"])
    changes, old_ff, new_ff = [], None, None
    if proj:
        p = db.scalar(select(Project).where(Project.code == proj["code"]))
        current = {t.activity_code: t for t in db.scalars(select(ScheduleTask).where(ScheduleTask.project_id == p.id))}
        for t in s["tasks"]:
            cur = current.get(t["activity_code"])
            if cur and cur.finish != t["finish"]:
                delta = (date.fromisoformat(t["finish"]) - date.fromisoformat(cur.finish)).days
                changes.append(dict(uid=t["uid"], name=t["name"], activity=t["activity_code"], old_finish=cur.finish, new_finish=t["finish"], delta_days=delta))
        old_ff = p.forecast_finish_period
        new_ff = max(t["finish"] for t in s["tasks"])[:7] if s["tasks"] else None
    unmapped = [t for t in s["tasks"] if not t["activity_code"]]
    checks = [
        dict(code="project_identified", ok=proj is not None, params=dict(code=proj["code"] if proj else None)),
        dict(code="rows_mapped", ok=not unmapped, params=dict(mapped=len(s["tasks"]) - len(unmapped), total=len(s["tasks"]))),
    ]
    return dict(kind="schedule_xml", engine="rules", ai_available=ai_enabled(), project=proj, title=s["title"],
                status_date=s["status_date"], tasks=s["tasks"], changes=changes,
                forecast_finish=dict(old=old_ff, new=new_ff), checks=checks, ready=proj is not None)


# ------------------------------------------------------------------ confirm

def confirm(db: Session, job_id: int, options: dict | None = None) -> dict:
    options = options or {}
    job = db.get(IngestionJob, job_id)
    if not job:
        raise IngestError("Job not found")
    if job.status != "preview":
        raise IngestError(f"Job is already {job.status}")
    pv = dict(job.preview)
    for item in pv.get("row_mapping", []):
        if item["label"] in options.get("row_mapping", {}):
            item["code"] = options["row_mapping"][item["label"]] or None
            item["method"] = "manual"
    path = DATA_DIR / job.stored_path
    handler = {"cost_workbook": _import_cost_workbook, "legacy_workbook": _import_legacy,
               "change_order_pdf": _import_change_order, "document_pdf": _import_document_pdf,
               "schedule_xml": _import_schedule}.get(pv["kind"])
    if not handler:
        raise IngestError("This file cannot be imported")
    result = handler(db, job, path, pv, options)
    job.status = "imported"
    job.result = result
    db.commit()
    invalidate()
    return result


def _project(db: Session, pv: dict, options: dict) -> Project:
    code = options.get("project_code") or (pv.get("project") or {}).get("code")
    p = db.scalar(select(Project).where(Project.code == code)) if code else None
    if not p:
        raise IngestError("Select the project this file belongs to")
    return p


def _add_document(db, project_id, doc_type, job, path, lang, pages=None, period=None) -> Document:
    d = Document(project_id=project_id, doc_type=doc_type, title=job.filename, filename=job.filename, rel_path=_rel(path),
                 language=lang, pages=len(pages) if pages else None, period=period, origin="upload")
    db.add(d)
    db.flush()
    for i, text in enumerate(pages or [], start=1):
        db.add(DocumentChunk(document_id=d.id, page=i, text=text))
    return d


def _save_template(db, pv: dict, act_map: dict, ct_map: dict, country: str | None, options: dict) -> dict:
    if pv.get("template"):
        tpl = db.get(MappingTemplate, pv["template"]["id"])
        tpl.times_used += 1
        tpl.last_used_at = datetime.utcnow()
        return dict(id=tpl.id, name=tpl.name, times_used=tpl.times_used, created=False)
    if not pv.get("fingerprint"):
        return {}
    tpl = MappingTemplate(
        name=options.get("template_name") or f"Formato importado {datetime.utcnow():%Y-%m-%d} ({pv['filename']})",
        fingerprint=pv["fingerprint"], doc_kind=pv["kind"], country_code=country,
        mapping=dict(activities={norm(l): c for l, c in act_map.items() if c}, cost_types={norm(l): c for l, c in ct_map.items() if c}),
        times_used=1, last_used_at=datetime.utcnow(),
    )
    db.add(tpl)
    db.flush()
    return dict(id=tpl.id, name=tpl.name, times_used=1, created=True)


def _import_cost_workbook(db, job, path, pv, options):
    p = _project(db, pv, options)
    before = kpi_snapshot(db, p.id)
    parsed = parse_workbook(path)
    act_map = {m["label"]: m["code"] for m in pv["row_mapping"]}
    ct_map = {m["label"]: m["code"] for m in pv["cost_type_mapping"]}
    doc = _add_document(db, p.id, "cost_workbook", job, path, COUNTRY_LANG.get(p.country_code, "es"))
    for model in (CostActual, Progress, Production):
        db.execute(delete(model).where(model.project_id == p.id))
    db.execute(delete(BudgetLine).where(BudgetLine.project_id == p.id, BudgetLine.origin == "original"))
    n_cost = n_budget = n_prog = 0
    for r in parsed["cost_rows"]:
        a = act_map.get(r["label"])
        if not a:
            continue
        ct = ct_map.get(r["cost_type_label"]) or "NA"
        if r["budget"] is not None:
            db.add(BudgetLine(project_id=p.id, activity_code=a, cost_type=ct, amount=r["budget"], origin="original", document_id=doc.id, locator=r["budget_loc"]))
            n_budget += 1
        for per, v in r["values"].items():
            if v:
                db.add(CostActual(project_id=p.id, period=per, activity_code=a, cost_type=ct, amount=v, document_id=doc.id, locator=r["locs"][per]))
                n_cost += 1
    scale = 100.0 if any(v > 1.5 for r in parsed["progress_rows"] for v in r["values"].values()) else 1.0
    series = defaultdict(dict)
    for r in parsed["progress_rows"]:
        a = act_map.get(r["label"])
        if a:
            for per, v in r["values"].items():
                series[(a, per)][r["series"]] = (v / scale, r["locs"][per])
    for (a, per), s in series.items():
        planned = s.get("planned", (0.0, None))[0]
        actual, loc = s.get("actual", (None, None))
        db.add(Progress(project_id=p.id, period=per, activity_code=a, planned_pct=planned, actual_pct=actual, document_id=doc.id, locator=loc))
        n_prog += 1
    for r in parsed["production_rows"]:
        db.add(Production(project_id=p.id, period=r["period"], km_welded_cum=r["km"] or 0, welds_cum=int(r["welds"] or 0),
                          weld_repairs_cum=int(r["repairs"] or 0), document_id=doc.id, locator=r["locator"]))
    tpl = _save_template(db, pv, act_map, ct_map, p.country_code, options)
    db.flush()
    refresh_kpis(db, [p.id])
    return dict(kind="cost_workbook", project_code=p.code, records=dict(cost_values=n_cost, budget_lines=n_budget, progress_points=n_prog,
                production_points=len(parsed["production_rows"])), before=before, after=kpi_snapshot(db, p.id), template=tpl, document_id=doc.id)


def _import_legacy(db, job, path, pv, options):
    np_ = dict(pv["new_project"])
    np_.update({k: v for k, v in options.get("project", {}).items() if v not in (None, "")})
    if not np_.get("code") or db.scalar(select(Project).where(Project.code == np_["code"])):
        raise IngestError("Project code is missing or already exists")
    parsed = parse_workbook(path)
    act_map = {m["label"]: m["code"] for m in pv["row_mapping"]}
    before_n = db.query(Project).filter(Project.status == "closed", Project.kind == "pipeline").count()
    finish = np_["finish_period"]
    desc = (
        f"Proyecto histórico importado desde la liquidación final ({job.filename}).",
        f"Historical project imported from its final cost settlement ({job.filename}).",
        f"Projeto histórico importado da liquidação final ({job.filename}).",
    )
    p = Project(code=np_["code"], name=np_["name"], country_code=np_["country"], client=np_.get("client") or "—",
                kind=np_["kind"], status="closed", contract_type="lump_sum", terrain=np_["terrain"], region=np_.get("region") or "",
                diameter_in=np_.get("diameter_in"), length_km=np_.get("length_km"), currency=np_["currency"],
                contract_value=np_["contract_value"], bid_margin_pct=round(np_["bid_margin_pct"], 4),
                start_period=np_["start_period"], planned_finish_period=finish, forecast_finish_period=finish,
                manager="—", description_es=desc[0], description_en=desc[1], description_pt=desc[2], origin="upload")
    db.add(p)
    db.flush()
    doc = _add_document(db, p.id, "legacy_workbook", job, path, COUNTRY_LANG.get(p.country_code, "es"), period=finish)
    meta = pv.get("meta", {})
    if "contract_value" in meta:
        p.contract_document_id, p.contract_locator = doc.id, meta["contract_value"]["locator"]
    p.bid_margin_document_id, p.bid_margin_locator = doc.id, meta.get("contract_value", {}).get("locator")
    acts = {}
    for r in parsed["legacy_rows"]:
        a = act_map.get(r["label"])
        if not a:
            continue
        acts.setdefault(a, r["actual_loc"])
        if r["budget"] is not None:
            db.add(BudgetLine(project_id=p.id, activity_code=a, cost_type="NA", amount=r["budget"], origin="original", document_id=doc.id, locator=r["budget_loc"]))
        if r["actual"] is not None:
            db.add(CostActual(project_id=p.id, period=finish, activity_code=a, cost_type="NA", amount=r["actual"], document_id=doc.id, locator=r["actual_loc"]))
    for a, loc in acts.items():
        db.add(Progress(project_id=p.id, period=finish, activity_code=a, planned_pct=1.0, actual_pct=1.0, document_id=doc.id, locator=loc))
    tpl = _save_template(db, pv, act_map, {}, p.country_code, options)
    db.flush()
    refresh_kpis(db, [p.id])
    after_n = db.query(Project).filter(Project.status == "closed", Project.kind == "pipeline").count()
    return dict(kind="legacy_workbook", project_code=p.code, created=True, records=dict(budget_lines=len(acts), cost_values=len(acts)),
                after=kpi_snapshot(db, p.id), template=tpl, benchmark_projects=dict(before=before_n, after=after_n), document_id=doc.id)


def _import_change_order(db, job, path, pv, options):
    p = _project(db, pv, options)
    before = kpi_snapshot(db, p.id)
    f = {x["field"]: x["value"] for x in pv["fields"]}
    f.update(options.get("fields", {}))
    lang = COUNTRY_LANG.get(p.country_code, "es")
    pages = read_pdf(path)
    doc = _add_document(db, p.id, "change_order", job, path, lang, pages, period=(f.get("submitted_date") or "")[:7] or None)
    title = f.get("title") or job.filename
    title_en = f.get("title_en") or title
    co = db.scalar(select(ChangeOrder).where(ChangeOrder.project_id == p.id, ChangeOrder.number == f.get("co_number")))
    created = co is None
    if created:
        co = ChangeOrder(project_id=p.id, number=f.get("co_number") or f"CO-{doc.id}")
        db.add(co)
    co.title_es = title if lang == "es" else title_en
    co.title_pt = title if lang == "pt" else title_en
    co.title_en = title_en
    co.cause = f.get("cause") or "scope"
    co.activity_code = f.get("activity_code") or "IND"
    co.amount = float(f.get("amount") or 0)
    co.status = f.get("status") or "pending"
    co.submitted_date = f.get("submitted_date") or datetime.utcnow().strftime("%Y-%m-%d")
    co.schedule_impact_days = int(f.get("schedule_impact_days") or 0)
    co.document_id, co.locator = doc.id, "p.1"
    db.flush()
    refresh_kpis(db, [p.id])
    return dict(kind="change_order_pdf", project_code=p.code, created=created, number=co.number,
                before=before, after=kpi_snapshot(db, p.id), document_id=doc.id)


def _import_document_pdf(db, job, path, pv, options):
    p = _project(db, pv, options)
    pages = read_pdf(path)
    doc = _add_document(db, p.id, pv.get("doc_type") or "other", job, path, COUNTRY_LANG.get(p.country_code, "es"), pages)
    return dict(kind="document_pdf", project_code=p.code, document_id=doc.id, pages=len(pages))


def _import_schedule(db, job, path, pv, options):
    p = _project(db, pv, options)
    before = kpi_snapshot(db, p.id)
    s = parse_mspdi(path)
    doc = _add_document(db, p.id, "schedule", job, path, COUNTRY_LANG.get(p.country_code, "es"))
    db.execute(delete(ScheduleTask).where(ScheduleTask.project_id == p.id))
    for t in s["tasks"]:
        db.add(ScheduleTask(project_id=p.id, uid=t["uid"], name=t["name"], activity_code=t["activity_code"],
                            baseline_start=t["baseline_start"] or t["start"], baseline_finish=t["baseline_finish"] or t["finish"],
                            start=t["start"], finish=t["finish"], pct_complete=t["pct"], document_id=doc.id, locator=f"Task UID {t['uid']}"))
    if s["tasks"]:
        p.forecast_finish_period = max(t["finish"] for t in s["tasks"])[:7]
    db.flush()
    refresh_kpis(db, [p.id])
    return dict(kind="schedule_xml", project_code=p.code, tasks=len(s["tasks"]), forecast_finish=p.forecast_finish_period,
                before=before, after=kpi_snapshot(db, p.id), document_id=doc.id)
